"""Ouverture d'applications et de fichiers."""
import configparser
import difflib
import os
import re
import subprocess
from pathlib import Path

DESKTOP_DIRS = [
    Path("/usr/share/applications"),
    Path("/usr/local/share/applications"),
    Path.home() / ".local/share/applications",
]

_EXEC_FIELD_CODE_RE = re.compile(r"%[a-zA-Z]")


def _iter_desktop_entries():
    for directory in DESKTOP_DIRS:
        if not directory.is_dir():
            continue
        for entry in directory.glob("*.desktop"):
            try:
                parser = configparser.ConfigParser(interpolation=None)
                parser.read(entry, encoding="utf-8")
                if "Desktop Entry" not in parser:
                    continue
                section = parser["Desktop Entry"]
                if section.get("NoDisplay", "false").lower() == "true":
                    continue
                if section.get("Type", "Application") != "Application":
                    continue
                name = section.get("Name")
                exec_cmd = section.get("Exec")
                if not name or not exec_cmd:
                    continue
                yield name, exec_cmd
            except Exception:
                continue


def _find_best_app(query: str, fuzzy: bool = True):
    query = query.strip().lower()
    entries = list(_iter_desktop_entries())
    names = [name for name, _ in entries]

    for name, exec_cmd in entries:
        if name.lower() == query:
            return name, exec_cmd

    substr_matches = [(n, e) for n, e in entries if query in n.lower()]
    if substr_matches:
        substr_matches.sort(key=lambda x: len(x[0]))
        return substr_matches[0]

    if not fuzzy:
        return None, None
    close = difflib.get_close_matches(query, names, n=1, cutoff=0.5)
    if close:
        for name, exec_cmd in entries:
            if name == close[0]:
                return name, exec_cmd
    return None, None


def open_application(name: str, fuzzy: bool = True) -> dict:
    """Lance une application installee dont le nom se rapproche de `name`."""
    app_name, exec_cmd = _find_best_app(name, fuzzy)
    if not exec_cmd:
        return {"success": False, "error": f"Aucune application trouvee pour '{name}'."}

    clean_cmd = _EXEC_FIELD_CODE_RE.sub("", exec_cmd).strip()
    try:
        subprocess.Popen(
            clean_cmd,
            shell=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            start_new_session=True,
        )
        return {"success": True, "app": app_name}
    except Exception as exc:
        return {"success": False, "error": str(exc)}


def open_path(path: str) -> dict:
    """Ouvre un fichier ou un dossier avec l'application par defaut (xdg-open)."""
    expanded = os.path.expanduser(os.path.expandvars(path))
    target = Path(expanded)
    if not target.exists():
        # tente une recherche simple dans le home si le chemin exact n'existe pas
        matches = list(Path.home().rglob(target.name)) if target.name else []
        if matches:
            target = matches[0]
        else:
            return {"success": False, "error": f"Chemin introuvable : {path}"}
    try:
        subprocess.Popen(
            ["xdg-open", str(target)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            start_new_session=True,
        )
        return {"success": True, "path": str(target)}
    except Exception as exc:
        return {"success": False, "error": str(exc)}


# Fenetres qu'on ne ferme jamais : le bureau, le terminal, Jarvis lui-meme, le shell.
_PROTECTED_CLASSES = ("nemo-desktop", "main.py", "terminal", "cinnamon", "mutter", "n/a")
_FOLDER_CLASSES = ("nemo.", "nautilus.", "thunar.", "dolphin.", "pcmanfm.", "caja.")


def _windows():
    out = subprocess.run(["wmctrl", "-lx"], capture_output=True, text=True, timeout=5).stdout
    for line in out.splitlines():
        parts = line.split(None, 4)
        if len(parts) == 5:
            yield parts[0], parts[2], parts[4]


def close_window(name: str = "", folders: bool = False) -> dict:
    """Ferme proprement (comme la croix de la fenetre) les fenetres dont le titre ou la classe contient `name`.
    `folders=True` : uniquement les fenetres de l'explorateur de fichiers (avec `name` = nom du dossier, optionnel)."""
    needle = (name or "").strip().lower()
    closed = []
    try:
        for wid, wclass, title in _windows():
            cls = wclass.lower()
            if any(p in cls for p in _PROTECTED_CLASSES):
                continue
            if folders:
                if not cls.startswith(_FOLDER_CLASSES):
                    continue
                if needle and needle not in title.lower():
                    continue
            else:
                squeezed = needle.replace(" ", "")          # « libre office » <-> « LibreOffice »
                haystack = (title + " " + wclass).lower()
                if not needle or (needle not in haystack and squeezed not in haystack.replace(" ", "")):
                    continue
            if subprocess.run(["wmctrl", "-ic", wid], timeout=5).returncode == 0:
                closed.append(title)
    except Exception as exc:
        return {"success": False, "error": str(exc)}
    if not closed:
        return {"success": False, "error": "Aucune fenetre correspondante ouverte."}
    return {"success": True, "closed": closed}


def create_spreadsheet(columns: int = 3, rows: int = 10, headers=None, title: str = "") -> dict:
    """Cree un tableau (.xlsx : en-tetes en gras, colonnes numerotees) dans ~/Documents et l'ouvre dans le tableur."""
    import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    try:
        columns, rows = int(columns), int(rows)
        if not (1 <= columns <= 50 and 1 <= rows <= 1000):
            return {"success": False, "error": "Taille de tableau invalide (1 a 50 colonnes, 1 a 1000 lignes)."}
        wb = Workbook()
        ws = wb.active
        ws.title = "Tableau"
        names = list(headers or [])[:columns] + [f"Colonne {i}" for i in range(len(headers or []) + 1, columns + 1)]
        side = Side(style="thin", color="999999")
        border = Border(left=side, right=side, top=side, bottom=side)
        for c, name in enumerate(names, 1):
            cell = ws.cell(row=1, column=c, value=name)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = 16
        for r in range(2, rows + 2):
            for c in range(1, columns + 1):
                ws.cell(row=r, column=c).border = border
        folder = Path.home() / "Documents"
        folder.mkdir(exist_ok=True)
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        base = re.sub(r"[^\w-]+", "_", title.strip()) if title.strip() else "tableau"
        path = folder / f"{base}_{stamp}.xlsx"
        wb.save(path)
    except Exception as exc:
        return {"success": False, "error": str(exc)}
    opened = open_path(str(path))
    return {"success": True, "path": str(path), "columns": columns, "rows": rows, "opened": opened.get("success", False)}
